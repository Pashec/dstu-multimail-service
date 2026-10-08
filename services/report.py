import os
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import DATA_DIR


def generate_excel_report(report_data: list[dict], summary_stats: dict) -> str:
    reports_dir = os.path.join(DATA_DIR, "reports")
    os.makedirs(reports_dir, exist_ok=True)

    timestamp_str = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"Отчет_рассылки_{timestamp_str}.xlsx"
    filepath = os.path.join(reports_dir, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Итоги рассылки"

    primary_blue = "003B95"
    success_green = "D1E7DD"
    success_text = "0F5132"
    error_red = "F8D7DA"
    error_text = "842029"
    warning_yellow = "FFF3CD"
    warning_text = "664D03"

    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_meta_lbl = Font(name="Calibri", size=11, bold=True, color="555555")
    font_meta_val = Font(name="Calibri", size=11, bold=True, color="111111")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)

    fill_title = PatternFill(start_color=primary_blue, end_color=primary_blue, fill_type="solid")
    fill_header = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CCCCCC")
    border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "ДГТУ РАССЫЛКА • ИТОГОВЫЙ ОТЧЕТ ОБ ОТПРАВКЕ ПИСЕМ"
    title_cell.font = font_title
    title_cell.fill = fill_title
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    total = summary_stats.get("total", len(report_data))
    sent = summary_stats.get("sent", sum(1 for r in report_data if r.get("status") == "Успешно"))
    errors_count = summary_stats.get("errors", sum(1 for r in report_data if r.get("status") == "Ошибка"))
    skipped = summary_stats.get("skipped", sum(1 for r in report_data if r.get("status") == "Пропущено"))
    subject = summary_stats.get("subject", "Не указана")

    meta_rows = [
        ("Дата и время формирования:", datetime.datetime.now().strftime("%d.%m.%Y %H:%M:%S")),
        ("Тема письма:", subject),
        ("Всего адресатов в списке:", total),
        ("Успешно доставлено:", f"{sent} ({(sent / total * 100):.1f}%)" if total > 0 else f"{sent}"),
        ("Ошибок при отправке:", errors_count),
        ("Пропущено строк:", skipped)
    ]

    curr_row = 3
    for label, val in meta_rows:
        ws.cell(row=curr_row, column=1, value=label).font = font_meta_lbl
        val_cell = ws.cell(row=curr_row, column=2, value=val)
        val_cell.font = font_meta_val
        curr_row += 1

    curr_row += 1

    headers = [
        "№",
        "Email адресата",
        "ФИО получателя",
        "Обращение",
        "Статус отправки",
        "Прикрепленные файлы",
        "Сведения об ошибке / Код SMTP",
        "Время фиксации"
    ]

    header_row_idx = curr_row
    ws.row_dimensions[header_row_idx].height = 28
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    curr_row += 1

    for item in report_data:
        ws.row_dimensions[curr_row].height = 20
        status = item.get("status", "—")

        row_vals = [
            item.get("index", curr_row - header_row_idx),
            item.get("email", ""),
            item.get("fio", ""),
            item.get("salutation", ""),
            status,
            item.get("personal_files", "—"),
            item.get("error_message", "—") if item.get("error_message") else "—",
            item.get("timestamp", "")
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = border_data

            if col_idx in (1, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (4, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if col_idx == 5:
                if status == "Успешно":
                    cell.fill = PatternFill(start_color=success_green, end_color=success_green, fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color=success_text)
                elif status == "Ошибка":
                    cell.fill = PatternFill(start_color=error_red, end_color=error_red, fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color=error_text)
                elif status == "Пропущено":
                    cell.fill = PatternFill(start_color=warning_yellow, end_color=warning_yellow, fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color=warning_text)

        curr_row += 1

    min_widths = [6, 26, 30, 16, 18, 30, 36, 20]
    for col_idx, min_w in enumerate(min_widths, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = min_w
        for row in range(header_row_idx, curr_row):
            val = ws.cell(row=row, column=col_idx).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)) + 3, 50))
        ws.column_dimensions[col_letter].width = max_len

    wb.save(filepath)
    return filepath
